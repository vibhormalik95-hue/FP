"""Build the course plan from its canonical Markdown, without duplicate prose.

Run with the supplied document runtime. The source is docs/135-hour-plan.md.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/135-hour-plan.md'
OUT = ROOT / 'deliverables/COMP9500-135-Hour-Plan.docx'


def inline(paragraph, text):
    for chunk in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not chunk:
            continue
        run = paragraph.add_run(chunk[2:-2] if chunk.startswith('**') else chunk[1:-1] if chunk.startswith('`') else chunk)
        if chunk.startswith('**'):
            run.bold = True
        if chunk.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def add_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', value) for value in row)]
    count = len(rows[0])
    widths = {2: [2.0, 4.9], 3: [1.5, 0.9, 4.5], 4: [0.5, 1.35, 0.55, 4.5], 5: [0.5, 0.9, 0.5, 2.0, 3.0]}.get(count, [6.9 / count] * count)
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, value in enumerate(values):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.02
            if widths[ci] <= 0.9:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if ri == 0:
                    run.bold = True
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                element = OxmlElement('w:' + side)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    element.set(qn('w:' + key), val)
                borders.append(element)
            props.append(borders)
            margin = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                element = OxmlElement('w:' + side)
                element.set(qn('w:w'), '85')
                element.set(qn('w:type'), 'dxa')
                margin.append(element)
            props.append(margin)
            if ri == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF2')
                props.append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    source = SOURCE.read_text()
    if '\u2014' in source:
        raise ValueError('The source contains an em dash.')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.06
    for name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            inline(p, line[2:].replace('-', ' '))
        elif line.startswith('### '):
            p = doc.add_paragraph(style='Heading 2')
            inline(p, line[4:])
        elif line.startswith('## '):
            title = line[3:]
            p = doc.add_paragraph(style='Heading 1')
            inline(p, title)
        else:
            p = doc.add_paragraph()
            inline(p, line.rstrip())
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('COMP 9500 prospective plan | ')
    run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'COMP 9500 prospective completion and verification plan'
    doc.core_properties.subject = '135 future active student hours from October to December 2026'
    doc.core_properties.author = 'Vibhor Malik; prepared with AI assistance'
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
